import docx
import sys

def extract_text(filename, output_file):
    doc = docx.Document(filename)
    with open(output_file, 'w', encoding='utf-8') as f:
        for para in doc.paragraphs:
            if para.text.strip():
                f.write(para.text + '\n')
        
        # also get tables
        for table in doc.tables:
            for row in table.rows:
                f.write(" | ".join([cell.text.replace("\n", " ") for cell in row.cells]) + '\n')
            f.write('\n')

if __name__ == '__main__':
    extract_text('G:/My Drive/RELS/Syriac/Lesson_05_Aphel_Stem.docx', 'G:/My Drive/RELS/Syriac/lesson5_dump.md')
